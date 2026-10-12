from pathlib import Path
import argparse
import sys

try:
    import numpy as np
    import pandas as pd
    import openpyxl  # Lector y escritor de archivos Excel.
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    print("Falta una biblioteca. Pega este comando en la terminal de VS Code:")
    prefijo = "& " if sys.platform == "win32" else ""
    print(f'{prefijo}"{sys.executable}" -m pip install pandas numpy openpyxl')
    print("Cuando termine la instalación, vuelve a ejecutar el archivo.")
    raise SystemExit(1)


COLUMNAS = [
    "AcelX1_SisGenEol", "AcelY1_SisGenEol", "AcelZ1_SisGenEol",
    "GiroX_SisGenEol", "GiroY_SisGenEol", "GiroZ_SisGenEol",
    "MagneX_SisGenEol", "MagneY_SisGenEol", "MagneZ_SisGenEol",
    "Ter_SisGenEol", "smp",
]


MODULO_CONTADOR = 65536


def es_contador(valor):
    return (
        np.isfinite(valor) and valor == np.floor(valor)
        and 0 <= valor < MODULO_CONTADOR
    )


def recuperar_estructura(original):
    if list(original.columns[:11]) != COLUMNAS:
        raise ValueError("Los primeros 11 encabezados no coinciden con el original esperado.")
    extras = list(original.columns[11:])
    if any(not str(c).startswith("Unnamed:") for c in extras):
        raise ValueError("Hay columnas adicionales con nombre. Deben revisarse antes de continuar.")

    numerico = original.replace(r"^\s*$", np.nan, regex=True).apply(
        pd.to_numeric, errors="raise"
    )
    valores = numerico.to_numpy(dtype=float)
    if np.isinf(valores).any():
        raise ValueError("Hay valores infinitos en el archivo original.")
    anchos = np.isfinite(valores).sum(axis=1)
    datos = numerico.iloc[:, :11].copy()
    cambios, celdas, fragmentos = [], [], []

    for i in np.flatnonzero(anchos > 11):
        ancho = int(anchos[i])
        fila = valores[i]
        if not (np.isfinite(fila[:ancho]).all() and np.isnan(fila[ancho:]).all()):
            raise ValueError(f"La fila {i + 2} tiene huecos internos y requiere revisión manual.")
        if not (12 <= ancho <= 21):
            raise ValueError(f"La fila {i + 2} tiene una longitud no contemplada: {ancho}.")
        if i == 0 or i + 1 == len(valores) or anchos[i - 1] != 11 or anchos[i + 1] != 11:
            raise ValueError(f"La fila {i + 2} no tiene dos filas vecinas completas para validarla.")

        inicio = ancho - 11
        bloque = fila[inicio:ancho].copy()
        anterior, contador, siguiente = valores[i - 1, 10], bloque[-1], valores[i + 1, 10]
        if not all(es_contador(v) for v in (anterior, contador, siguiente)):
            raise ValueError(f"No se puede validar el contador de la fila {i + 2}.")
        paso_anterior = int((contador - anterior) % MODULO_CONTADOR)
        paso_siguiente = int((siguiente - contador) % MODULO_CONTADOR)
        if paso_anterior not in (1, 2) or paso_siguiente != 1:
            raise ValueError(f"El bloque final de la fila {i + 2} no encaja con sus vecinos.")

        datos.iloc[i, :] = bloque
        cambios.append({
            "fila_excel": int(i + 2), "celdas_presentes": ancho,
            "columna_inicio_bloque": get_column_letter(inicio + 1),
            "columna_fin_bloque": get_column_letter(ancho),
            "smp_anterior": int(anterior), "smp_recuperado": int(contador),
            "smp_siguiente": int(siguiente), "paso_anterior_modulo": paso_anterior,
            "paso_siguiente_modulo": paso_siguiente, "valores_fragmento": inicio,
        })
        for j, variable in enumerate(COLUMNAS):
            origen = inicio + j
            celdas.append({
                "fila_excel": int(i + 2),
                "celda_origen": f"{get_column_letter(origen + 1)}{i + 2}",
                "variable_destino": variable, "valor_original_recuperado": float(fila[origen]),
            })
        for j in range(inicio):
            fragmentos.append({
                "fila_excel": int(i + 2),
                "celda_origen": f"{get_column_letter(j + 1)}{i + 2}",
                "valor_original": float(fila[j]),
                "motivo": "Fragmento previo al bloque completo; registro no reconstruible con certeza",
            })

    # Un valor perdido se conserva como perdido; aquí no se imputa.
    presentes = datos["smp"].dropna().to_numpy(dtype=float)
    if not all(es_contador(v) for v in presentes):
        raise ValueError("Quedan contadores no enteros o fuera del rango observado.")

    saltos, reinicios = [], []
    contador = datos["smp"].to_numpy(dtype=float)
    for i in range(1, len(contador)):
        previo, actual = contador[i - 1], contador[i]
        if not (np.isfinite(previo) and np.isfinite(actual)):
            continue
        if actual < previo:
            if previo != MODULO_CONTADOR - 1 or actual != 0:
                raise ValueError(f"Hay un descenso no explicado del contador en la fila {i + 2}.")
            reinicios.append(int(i + 2))
        paso = int((actual - previo) % MODULO_CONTADOR)
        if paso not in (1, 2):
            raise ValueError(f"Hay un salto no contemplado del contador en la fila {i + 2}: {paso}.")
        if paso == 2:
            saltos.append({
                "fila_excel_anterior": int(i + 1), "fila_excel_actual": int(i + 2),
                "smp_anterior": int(previo), "smp_actual": int(actual),
                "posiciones_sin_registro_completo": paso - 1,
            })

    return datos, cambios, celdas, fragmentos, saltos, reinicios


def rellenar_vecinos(serie, cantidad=2):
    """Promedia los valores válidos de las filas más cercanas.

    Solo utiliza valores originalmente presentes. Una estimación nunca se
    convierte en vecino de otra estimación. Los empates se resuelven dando
    prioridad a la fila de menor número.
    """
    if cantidad < 1:
        raise ValueError("La cantidad de vecinos debe ser al menos 1.")

    valores = serie.to_numpy(dtype=float, copy=True)
    resultado = valores.copy()
    validos = np.flatnonzero(~np.isnan(valores))
    vacios = np.flatnonzero(np.isnan(valores))
    cambios = []

    if len(vacios) and not len(validos):
        raise ValueError(f"La columna {serie.name} está completamente vacía.")

    for fila in vacios:
        centro = int(np.searchsorted(validos, fila))
        candidatos = validos[
            max(0, centro - cantidad): min(len(validos), centro + cantidad)
        ]
        vecinos = sorted(
            candidatos, key=lambda i: (abs(int(i) - int(fila)), int(i))
        )[:cantidad]
        estimacion = float(np.mean(valores[vecinos]))
        resultado[fila] = estimacion
        cambios.append({
            "fila_excel": int(fila) + 2,
            "columna": serie.name,
            "valor_estimado": estimacion,
            "metodo": "Promedio de vecinos por posición de fila",
            "filas_vecinas_excel": ", ".join(str(int(i) + 2) for i in vecinos),
            "valores_vecinos": ", ".join(str(valores[i]) for i in vecinos),
        })

    return pd.Series(resultado, index=serie.index, name=serie.name), cambios


def limpiar(original, cantidad_vecinos=2):
    """Devuelve datos, registro de estimaciones y filas para revisar."""
    faltantes = [nombre for nombre in COLUMNAS if nombre not in original.columns]
    if faltantes:
        raise ValueError(f"No encuentro estos encabezados en el Excel: {faltantes}")

    adicionales = [nombre for nombre in original.columns if nombre not in COLUMNAS]
    con_extras = (
        original[adicionales].notna().any(axis=1)
        if adicionales else pd.Series(False, index=original.index)
    )
    revisar = original.loc[con_extras].copy()
    revisar.insert(0, "fila_excel_original", revisar.index.to_numpy() + 2)

    # Esta selección elimina las columnas adicionales, sin eliminar filas.
    datos = original[COLUMNAS].copy()
    datos = datos.replace(r"^\s*$", np.nan, regex=True)
    for columna in COLUMNAS:
        # Un texto inesperado produce un mensaje para revisarlo.
        datos[columna] = pd.to_numeric(datos[columna], errors="raise")

    cambios = []
    for columna in COLUMNAS[:-1]:
        datos[columna], registro = rellenar_vecinos(
            datos[columna], cantidad_vecinos
        )
        cambios.extend(registro)

    # smp parece ser un contador de muestras. No se promedia como un sensor.
    # En el original entregado solo falta en la última fila.
    vacios_smp = np.flatnonzero(datos["smp"].isna().to_numpy())
    for fila in vacios_smp:
        if fila != len(datos) - 1 or fila < 2:
            raise ValueError(
                "Hay un vacío en smp fuera del final del archivo. "
                "Revisa el contador antes de completar ese dato."
            )
        penultimo = datos["smp"].iloc[fila - 1]
        antepenultimo = datos["smp"].iloc[fila - 2]
        if (
            pd.isna(penultimo) or pd.isna(antepenultimo)
            or not float(penultimo).is_integer()
            or penultimo != antepenultimo + 1
        ):
            raise ValueError("No puedo inferir el último smp: su secuencia no es clara.")
        estimacion = float(penultimo + 1)
        datos.loc[datos.index[fila], "smp"] = estimacion
        cambios.append({
            "fila_excel": int(fila) + 2,
            "columna": "smp",
            "valor_estimado": estimacion,
            "metodo": "Contador inferido: número anterior + 1",
            "filas_vecinas_excel": f"{int(fila)}, {int(fila) + 1}",
            "valores_vecinos": f"{antepenultimo}, {penultimo}",
        })

    columnas_registro = [
        "fila_excel", "columna", "valor_estimado", "metodo",
        "filas_vecinas_excel", "valores_vecinos",
    ]
    registro = pd.DataFrame(cambios, columns=columnas_registro)
    assert len(datos) == len(original), "Cambió la cantidad de filas."
    assert list(datos.columns) == COLUMNAS, "Cambió el orden de las columnas."
    return datos, registro, revisar, adicionales


def main():
    argumentos = argparse.ArgumentParser(description="Limpieza inicial del Excel eólico")
    argumentos.add_argument("archivo", nargs="?", help="Ruta opcional del Excel original")
    argumentos.add_argument("--salida", help="Ruta opcional del resultado")
    argumentos.add_argument("--vecinos", type=int, default=2, help="Vecinos por vacío (2 por defecto)")
    opciones = argumentos.parse_args()

    carpeta = Path(__file__).resolve().parent
    if opciones.archivo:
        entrada = Path(opciones.archivo).resolve()
    else:
        nombres = ["CF-2023-I-204-1230_DAT.xlsx", "CF-2023-I-204-1230_DAT(1).xlsx", "CF-2023-I-204-1230_DAT(2).xlsx"]
        entrada = next((carpeta / n for n in nombres if (carpeta / n).is_file()), None)
        if entrada is None:
            raise FileNotFoundError(
                "Coloca limpieza-de-dataset-2.py en la misma carpeta que tu Excel original. "
                "Nombre esperado: CF-2023-I-204-1230_DAT.xlsx"
            )

    salida = (
        Path(opciones.salida).resolve() if opciones.salida
        else entrada.parent / "Dataset_eolico_filtrado_2.xlsx"
    )
    if salida.resolve() == entrada.resolve():
        raise ValueError("El resultado debe tener un nombre distinto del original.")

    print("Leyendo el Excel original...", flush=True)
    original = pd.read_excel(
        entrada, sheet_name=0, engine="openpyxl", na_values=["NULL", "null", "None", " "]
    )
    realineados, correcciones, celdas, fragmentos, saltos, reinicios = recuperar_estructura(original)
    datos, registro, _, _ = limpiar(realineados, opciones.vecinos)
    adicionales = list(original.columns[11:])
    revisar = original.loc[original[adicionales].notna().any(axis=1)].copy() if adicionales else original.iloc[:0].copy()
    revisar.insert(0, "fila_excel_original", revisar.index + 2)
    print("Guardando la nueva copia...", flush=True)
    salida.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        datos.to_excel(escritor, sheet_name="Datos", index=False)
        registro.to_excel(escritor, sheet_name="Valores_estimados", index=False)
        pd.DataFrame(correcciones).to_excel(escritor, sheet_name="Filas_realineadas", index=False)
        pd.DataFrame(celdas).to_excel(escritor, sheet_name="Origen_valores", index=False)
        pd.DataFrame(fragmentos).to_excel(escritor, sheet_name="Fragmentos_conservados", index=False)
        pd.DataFrame(saltos).to_excel(escritor, sheet_name="Saltos_contador", index=False)
        revisar.to_excel(escritor, sheet_name="Filas_para_revisar", index=False)
        for hoja in escritor.book.worksheets:
            hoja.freeze_panes = "A2"

    print("\nLIMPIEZA INICIAL COMPLETADA")
    print(f"Filas conservadas: {len(datos)}")
    print(f"Columnas conservadas: {len(datos.columns)}")
    print(f"Columnas adicionales eliminadas de Datos: {len(adicionales)}")
    print(f"Vacíos completados: {len(registro)}")
    print(f"Vacíos restantes: {int(datos.isna().sum().sum())}")
    print(f"Filas con valores en columnas extra, para revisar: {len(revisar)}")
    print(f"Filas realineadas y validadas con el contador: {len(correcciones)}")
    print("Los fragmentos restantes se conservaron en una hoja aparte.")
    print(f"Archivo guardado en: {salida}")
    print("\nÚltimas filas:")
    print(datos.tail(3).to_string(index=False))


if __name__ == "__main__":
    try:
        main()
    except PermissionError:
        print("Cierra el Excel de salida si está abierto y ejecuta de nuevo el código.")
        raise SystemExit(1)
