"""Genera resultado_analisis_2.xlsx con pandas y openpyxl.

Guarda este archivo junto a Dataset_eolico_filtrado_2.xlsx y al Excel original.
Abre el archivo .py en VS Code y pulsa el botón de ejecutar.
Las estadísticas se calculan en Python; ejecuta de nuevo para actualizarlas.
"""

from pathlib import Path
import argparse
import math
import sys

try:
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as error:
    print(f"Falta la biblioteca: {error.name}")
    print("Pega esta línea en la terminal y presiona Enter:")
    prefijo = "& " if sys.platform == "win32" else ""
    print(f'{prefijo}"{sys.executable}" -m pip install pandas openpyxl')
    print("Espera a que termine; luego vuelve a ejecutar este archivo.")
    raise SystemExit(1)

COLUMNAS = [
    "AcelX1_SisGenEol", "AcelY1_SisGenEol", "AcelZ1_SisGenEol",
    "GiroX_SisGenEol", "GiroY_SisGenEol", "GiroZ_SisGenEol",
    "MagneX_SisGenEol", "MagneY_SisGenEol", "MagneZ_SisGenEol",
    "Ter_SisGenEol", "smp",
]
SENSORES = COLUMNAS[:-1]


def calcular_estadisticas(datos, factor_iqr=1.5):
    """Usa cada columna sin sus vacíos; no elimina filas ni valores atípicos."""
    registros = []
    for columna in COLUMNAS:
        serie = datos[columna].dropna()
        n = len(serie)
        distintos = int(serie.nunique())
        q1 = serie.quantile(0.25) if n else float("nan")
        q3 = serie.quantile(0.75) if n else float("nan")
        iqr = q3 - q1
        desviacion = serie.std(ddof=1) if n >= 2 else float("nan")
        varianza = serie.var(ddof=1) if n >= 2 else float("nan")
        poblacional = serie.std(ddof=0) if n else float("nan")
        if distintos == 1:
            poblacional = 0.0
            if n >= 2:
                desviacion = varianza = 0.0
        atipicos = int(((serie < q1 - factor_iqr * iqr) |
                        (serie > q3 + factor_iqr * iqr)).sum())
        registros.append({
            "Variable": columna, "N": n,
            "Promedio": serie.mean() if n else float("nan"),
            "Mediana": serie.median() if n else float("nan"),
            "Desv. estándar muestral": desviacion,
            "Varianza muestral": varianza,
            "Mínimo": serie.min() if n else float("nan"),
            "Q1 (25%)": q1, "Q3 (75%)": q3,
            "Máximo": serie.max() if n else float("nan"),
            "Rango": serie.max() - serie.min() if n else float("nan"),
            "Vacíos": int(datos[columna].isna().sum()),
            "Vacíos (%)": float(datos[columna].isna().mean()),
            "Valores distintos": distintos,
            "Desv. estándar poblacional": poblacional, "IQR": iqr,
            "Asimetría": serie.skew() if n >= 3 and poblacional > 0 else float("nan"),
            "Curtosis (exceso)": serie.kurt() if n >= 4 and poblacional > 0 else float("nan"),
            "Atípicos IQR": atipicos,
            "Atípicos (%)": atipicos / n if n else float("nan"),
            "Sin variación": "Sí" if n and distintos == 1 else "No" if n else "n.a.",
        })
    return pd.DataFrame(registros)


def dar_formato(hoja, tabla, fila_cabecera=1, filtro=True):
    """Formatea encabezados y columnas sin recorrer todas las celdas de datos."""
    hoja.sheet_view.showGridLines = False
    hoja.freeze_panes = f"B{fila_cabecera + 1}"
    hoja.row_dimensions[fila_cabecera].height = 44
    for i, nombre in enumerate(tabla.columns, start=1):
        celda = hoja.cell(fila_cabecera, i)
        celda.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="29465B")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        dimension = hoja.column_dimensions[get_column_letter(i)]
        dimension.width = 29 if i == 1 else 19
        dimension.number_format = "0.000000"
        if nombre in {"N", "Vacíos", "Valores distintos", "Atípicos IQR",
                      "Valores nulos", "Valores no nulos", "Cantidad de nulos",
                      "Fila de Excel", "fila_excel"}:
            dimension.number_format = "#,##0"
        elif nombre in {"Vacíos (%)", "Atípicos (%)"} or "Porcentaje" in str(nombre):
            dimension.number_format = "0.0000%"
        elif "Varianza" in str(nombre):
            dimension.number_format = "0.000000E+00"
        if len(tabla) <= 1000:
            for fila in range(fila_cabecera + 1, fila_cabecera + len(tabla) + 1):
                hoja.cell(fila, i).number_format = dimension.number_format
    if filtro:
        hoja.auto_filter.ref = (
            f"A{fila_cabecera}:{get_column_letter(len(tabla.columns))}"
            f"{fila_cabecera + len(tabla)}"
        )


def crear_resultado(entrada, original, salida, factor_iqr=1.5):
    print("1/3: Leyendo el dataset limpio...", flush=True)
    with pd.ExcelFile(entrada, engine="openpyxl") as libro:
        hoja = "Datos" if "Datos" in libro.sheet_names else libro.sheet_names[0]
        datos = pd.read_excel(libro, sheet_name=hoja)
        registro = (pd.read_excel(libro, sheet_name="Valores_estimados")
                    if "Valores_estimados" in libro.sheet_names else pd.DataFrame())
    faltantes = [c for c in COLUMNAS if c not in datos.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas en el dataset limpio: {faltantes}")
    datos = datos[COLUMNAS].replace(r"^\s*$", float("nan"), regex=True).copy()
    for columna in COLUMNAS:
        datos[columna] = pd.to_numeric(datos[columna], errors="raise")
    if datos.empty:
        raise ValueError("El dataset no contiene filas.")
    if datos.isin([float("inf"), float("-inf")]).any().any():
        raise ValueError("Hay valores infinitos; revisa el dataset antes del análisis.")

    print("2/3: Revisando el original y calculando estadísticas...", flush=True)
    bruto = pd.read_excel(original, engine="openpyxl", sheet_name=0,
                          na_values=["NULL", "null", "None", " "])
    bruto = bruto.replace(r"^\s*$", float("nan"), regex=True)
    if any(c not in bruto.columns for c in COLUMNAS):
        raise ValueError("El Excel original no tiene los 11 encabezados esperados.")
    filas_nulas = bruto.loc[bruto[COLUMNAS].isna().any(axis=1), COLUMNAS].copy()
    filas_nulas.insert(0, "Fila de Excel", filas_nulas.index + 2)
    adicionales = [c for c in bruto.columns if c not in COLUMNAS]
    mascara = (bruto[adicionales].notna().any(axis=1) if adicionales
               else pd.Series(False, index=bruto.index))
    extras = bruto.loc[mascara].copy()
    extras.insert(0, "Fila de Excel", extras.index + 2)
    if registro.empty:
        registro = pd.DataFrame(columns=["fila_excel", "columna", "valor_estimado",
                                         "metodo", "filas_vecinas_excel", "valores_vecinos"])
    estadisticas = calcular_estadisticas(datos, factor_iqr)
    basicas = estadisticas.iloc[:, :11]
    distribucion = estadisticas[["Variable"] + list(estadisticas.columns[11:])]
    correlaciones = datos[SENSORES].corr(method="pearson")
    tablas = {
        "Datos principales": datos,
        "Tipos de datos": pd.DataFrame({
            "Variable": COLUMNAS, "Tipo en pandas": datos.dtypes.astype(str).to_numpy(),
            "Clasificación": ["Integer" if pd.api.types.is_integer_dtype(datos[c])
                              else "Float" for c in COLUMNAS],
            "Uso": ["Sensor"] * 10 + ["Contador de muestras"],
        }),
        "Resumen de nulos": pd.DataFrame({
            "Columna": bruto.columns, "Valores nulos": bruto.isna().sum().to_numpy(),
            "Valores no nulos": bruto.notna().sum().to_numpy(),
            "Porcentaje de nulos": bruto.isna().mean().to_numpy(),
        }),
        "Nulos principales": pd.DataFrame({
            "Variable": COLUMNAS, "Cantidad de nulos": datos.isna().sum().to_numpy(),
            "Porcentaje de nulos": datos.isna().mean().to_numpy(),
        }),
        "Filas con nulos": filas_nulas, "Datos adicionales": extras,
        "Valores estimados": registro,
    }
    print("3/3: Guardando el Excel; espera hasta ver ANÁLISIS COMPLETADO...", flush=True)
    salida.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        basicas.to_excel(escritor, sheet_name="Estadisticas", index=False, startrow=5)
        distribucion.to_excel(escritor, sheet_name="Estadisticas", index=False, startrow=20)
        correlaciones.to_excel(escritor, sheet_name="Correlaciones", index_label="Sensor")
        for nombre, tabla in tablas.items():
            tabla.to_excel(escritor, sheet_name=nombre, index=False)
            dar_formato(escritor.sheets[nombre], tabla)
        est = escritor.sheets["Estadisticas"]
        dar_formato(est, basicas, fila_cabecera=6, filtro=False)
        dar_formato(est, distribucion, fila_cabecera=21, filtro=False)
        est.freeze_panes = "B7"
        # Los dos bloques comparten columnas: se formatea cada celda del resumen.
        for tabla, encabezado in [(basicas, 6), (distribucion, 21)]:
            for j, nombre in enumerate(tabla.columns, start=1):
                for fila in range(encabezado + 1, encabezado + len(tabla) + 1):
                    celda = est.cell(fila, j)
                    celda.number_format = ("0.0000%" if nombre in {"Vacíos (%)", "Atípicos (%)"} else
                                           "0.000000E+00" if "Varianza" in nombre else
                                           "#,##0" if nombre in {"N", "Vacíos", "Valores distintos", "Atípicos IQR"}
                                           else "0.000000")
                    if nombre == "Desv. estándar muestral":
                        celda.font = Font(bold=True, color="29465B")
                        celda.fill = PatternFill("solid", fgColor="E9F1F7")
        textos = {
            1: "Estadísticas del dataset eólico",
            2: f"Datos: {entrada.name} | Original: {original.name}",
            3: f"Filas originales: {len(bruto)} | Filas analizadas: {len(datos)} | Variables: 11",
            4: f"Filas con datos adicionales en el original: {len(extras)}. Se incluyen; revisa Datos adicionales antes del PCA.",
            35: "Cómo leer los resultados",
            36: "N: valores presentes. Promedio y mediana: medidas del centro de los datos.",
            37: "Desviación estándar muestral: dispersión respecto del promedio, con divisor N−1. Poblacional: divisor N.",
            38: "Varianza muestral: cuadrado de la desviación estándar muestral. Rango: máximo menos mínimo.",
            39: f"Q1 y Q3: percentiles 25 y 75. IQR: Q3−Q1. Atípicos: fuera de [Q1−{factor_iqr}×IQR, Q3+{factor_iqr}×IQR].",
            40: "Los atípicos se cuentan y se conservan. Asimetría: dirección de la cola. Curtosis: exceso respecto de la normal.",
            41: "smp es el contador de muestras; se excluye de la correlación entre los diez sensores.",
            42: "Resumen de nulos, Filas con nulos y Datos adicionales describen el original; Nulos principales describe el limpio.",
            43: "Las cifras se calculan en Python. Si modificas los datos, ejecuta este archivo otra vez para actualizarlas.",
        }
        for fila, texto in textos.items():
            est.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=11)
            celda = est.cell(fila, 1, texto)
            celda.font = Font(name="Arial", size=14 if fila == 1 else 10,
                              bold=fila in {1, 35}, color="29465B")
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            est.row_dimensions[fila].height = 32 if fila in {1, 4} else 28
        cor = escritor.sheets["Correlaciones"]
        dar_formato(cor, correlaciones.reset_index(names="Sensor"), filtro=False)
        for fila in cor.iter_rows(min_row=2, max_row=11, min_col=2, max_col=11):
            for celda in fila:
                celda.number_format = "0.000"
                celda.alignment = Alignment(horizontal="center")
        cor.conditional_formatting.add("B2:K11", ColorScaleRule(
            start_type="num", start_value=-1, start_color="A9C6DE",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=1, end_color="EABFA9"))
        escritor.sheets["Valores estimados"].column_dimensions["D"].width = 48
        escritor.sheets["Datos principales"].freeze_panes = "A2"
        escritor.sheets["Datos principales"].column_dimensions["K"].number_format = "0.######"

    print("\nANÁLISIS COMPLETADO")
    print(f"Filas analizadas: {len(datos)} | Variables: 11")
    print(f"Valores nulos restantes: {int(datos.isna().sum().sum())}")
    print("\nDesviaciones estándar muestrales:")
    print(estadisticas.set_index("Variable")["Desv. estándar muestral"].to_string())
    print(f"\nArchivo guardado en:\n{salida}")
    return estadisticas


def main():
    parser = argparse.ArgumentParser(description="Análisis estadístico de la segunda limpieza")
    parser.add_argument("archivo", nargs="?", help="Ruta opcional del dataset limpio")
    parser.add_argument("--original", help="Ruta opcional del Excel original")
    parser.add_argument("--salida", help="Ruta opcional del Excel de resultados")
    parser.add_argument("--factor-iqr", type=float, default=1.5)
    args = parser.parse_args()
    if not math.isfinite(args.factor_iqr) or args.factor_iqr <= 0:
        raise ValueError("El factor IQR debe ser positivo y finito.")
    carpeta = Path(__file__).resolve().parent
    entrada = Path(args.archivo).resolve() if args.archivo else carpeta / "Dataset_eolico_filtrado_2.xlsx"
    if not entrada.is_file():
        raise FileNotFoundError("Coloca este código junto a Dataset_eolico_filtrado_2.xlsx; primero ejecuta la limpieza.")
    original = (Path(args.original).resolve() if args.original else next(
        (entrada.parent / n for n in ["CF-2023-I-204-1230_DAT.xlsx", "CF-2023-I-204-1230_DAT(1).xlsx"]
         if (entrada.parent / n).is_file()), None))
    if original is None or not original.is_file():
        raise FileNotFoundError("Coloca también el Excel original CF-2023-I-204-1230_DAT.xlsx en esa carpeta.")
    salida = Path(args.salida).resolve() if args.salida else entrada.parent / "resultado_analisis_2.xlsx"
    if salida.resolve() in {entrada.resolve(), original.resolve()}:
        raise ValueError("El resultado debe tener un nombre diferente de los archivos de entrada.")
    crear_resultado(entrada, original, salida, args.factor_iqr)


if __name__ == "__main__":
    try:
        main()
    except PermissionError:
        print("Cierra los archivos Excel que estén abiertos y vuelve a ejecutar este código.")
        raise SystemExit(1)
    except (FileNotFoundError, ValueError) as error:
        print(f"\n{error}")
        raise SystemExit(1)
