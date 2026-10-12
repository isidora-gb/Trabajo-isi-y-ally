# -*- coding: utf-8 -*-
"""
TAREA 2 - Tablas comparativas SIMPLES ANTES / DESPUÉS de la limpieza inicial.

Requisitos (terminal de VS Code):
    python -m pip install pandas numpy openpyxl

Uso (con un Excel original y uno filtrado en la misma carpeta):
    python Tablas_comparativas_Tarea2_SIMPLES.py

Uso cuando hay varias copias de los archivos:
    python Tablas_comparativas_Tarea2_SIMPLES.py --original "CF-2023-I-204-1230_DAT(3).xlsx" --limpio "Dataset_eolico_filtrado_2(3).xlsx"

Este programa NO modifica ninguno de los dos Excel de entrada.
Solo crea un Excel NUEVO con seis hojas y estadísticos de los 10 sensores.
El contador smp solo se evalúa para nulos y control de registros; NO se
considera una señal física y se excluye del análisis de atípicos.
"""

from pathlib import Path
import argparse
import sys

try:
    import numpy as np
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as error:
    print(f"Falta instalar: {error.name}")
    print(f'Instala las bibliotecas con: "{sys.executable}" -m pip install pandas numpy openpyxl')
    raise SystemExit(1)

SENSORES = [
    "AcelX1_SisGenEol", "AcelY1_SisGenEol", "AcelZ1_SisGenEol",
    "GiroX_SisGenEol", "GiroY_SisGenEol", "GiroZ_SisGenEol",
    "MagneX_SisGenEol", "MagneY_SisGenEol", "MagneZ_SisGenEol",
    "Ter_SisGenEol",
]
COLUMNAS = SENSORES + ["smp"]
FACTOR_IQR = 1.5


def localizar_archivo(carpeta, patron, argumento, descripcion):
    if argumento:
        ruta = Path(argumento).expanduser()
        if not ruta.is_absolute():
            ruta = carpeta / ruta
        ruta = ruta.resolve()
        if not ruta.is_file():
            raise FileNotFoundError(f"No existe el Excel {descripcion}: {ruta}")
        return ruta

    opciones = sorted(p for p in carpeta.glob(patron) if not p.name.startswith("~$"))
    if len(opciones) != 1:
        encontrados = [p.name for p in opciones]
        raise ValueError(
            f"Necesito identificar un único Excel {descripcion}. Encontrados: {encontrados}. "
            f"Si tienes varias versiones, ejecuta con --original y --limpio."
        )
    return opciones[0]


def leer_datos(ruta, limpio=False):
    hoja = "Datos" if limpio else 0
    df = pd.read_excel(ruta, sheet_name=hoja, engine="openpyxl", na_values=["NULL", "null", "None", " "])
    faltantes = [c for c in COLUMNAS if c not in df.columns]
    if faltantes:
        raise ValueError(f"Faltan variables en {ruta.name}: {faltantes}")
    datos = df[COLUMNAS].copy()
    for col in COLUMNAS:
        datos[col] = pd.to_numeric(datos[col], errors="raise")
    if not np.isfinite(datos.dropna().to_numpy(dtype=float)).all():
        raise ValueError(f"El Excel {ruta.name} contiene valores infinitos.")
    return df, datos


def estadisticas(datos):
    filas = []
    for sensor in SENSORES:
        s = datos[sensor].dropna()
        if len(s) < 2:
            raise ValueError(f"No hay suficientes registros para {sensor}")
        q1, q3 = s.quantile([0.25, 0.75])
        iqr = q3 - q1
        li = q1 - FACTOR_IQR * iqr
        ls = q3 + FACTOR_IQR * iqr
        num_atipicos = int(((s < li) | (s > ls)).sum())
        filas.append({
            "Sensor": sensor.replace("_SisGenEol", ""),
            "N": int(s.count()),
            "Media": float(s.mean()),
            "Mediana": float(s.median()),
            "Desviación estándar": float(s.std(ddof=1)),
            "Mínimo": float(s.min()),
            "Máximo": float(s.max()),
            "Q1": float(q1),
            "Q3": float(q3),
            "IQR": float(iqr),
            "Límite inferior IQR": float(li),
            "Límite superior IQR": float(ls),
            "Atípicos IQR": num_atipicos,
            "Atípicos (%)": 100 * num_atipicos / len(s),
        })
    return pd.DataFrame(filas).set_index("Sensor")


def preparar_tablas(bruto, original, limpio_full, limpio):
    antes = estadisticas(original)
    despues = estadisticas(limpio)

    resumen = pd.DataFrame({
        "Sensor": antes.index,
        "Nulos antes": [int(original[s + "_SisGenEol"].isna().sum()) for s in antes.index],
        "Nulos después": [int(limpio[s + "_SisGenEol"].isna().sum()) for s in antes.index],
        "DS antes": antes["Desviación estándar"].to_numpy(),
        "DS después": despues["Desviación estándar"].to_numpy(),
        "Reducción DS (%)": (100 * (antes["Desviación estándar"] - despues["Desviación estándar"]) / antes["Desviación estándar"].replace(0, np.nan)).to_numpy(),
        "Atípicos antes": antes["Atípicos IQR"].to_numpy(),
        "Atípicos después": despues["Atípicos IQR"].to_numpy(),
    })

    tabla_est = pd.concat([
        antes.assign(Etapa="Antes"),
        despues.assign(Etapa="Después"),
    ]).reset_index()
    tabla_est = tabla_est[["Etapa", "Sensor", "N", "Media", "Mediana", "Desviación estándar", "Mínimo", "Máximo", "Q1", "Q3", "IQR"]]

    tabla_nulos = pd.DataFrame({
        "Variable": COLUMNAS,
        "Nulos antes": [int(original[c].isna().sum()) for c in COLUMNAS],
        "Nulos después": [int(limpio[c].isna().sum()) for c in COLUMNAS],
        "Nulos antes (%)": [100 * original[c].isna().mean() for c in COLUMNAS],
        "Nulos después (%)": [100 * limpio[c].isna().mean() for c in COLUMNAS],
    })

    tabla_atipicos = pd.DataFrame({
        "Sensor": antes.index,
        "LI antes": antes["Límite inferior IQR"].to_numpy(),
        "LS antes": antes["Límite superior IQR"].to_numpy(),
        "Atípicos antes": antes["Atípicos IQR"].to_numpy(),
        "Atípicos antes (%)": antes["Atípicos (%)"].to_numpy(),
        "LI después": despues["Límite inferior IQR"].to_numpy(),
        "LS después": despues["Límite superior IQR"].to_numpy(),
        "Atípicos después": despues["Atípicos IQR"].to_numpy(),
        "Atípicos después (%)": despues["Atípicos (%)"].to_numpy(),
    })

    adicionales = [c for c in bruto.columns if c not in COLUMNAS]
    filas_extra_original = int(bruto[adicionales].notna().any(axis=1).sum()) if adicionales else 0
    adicionales_limpio = [c for c in limpio_full.columns if c not in COLUMNAS]
    filas_extra_limpio = int(limpio_full[adicionales_limpio].notna().any(axis=1).sum()) if adicionales_limpio else 0

    tabla_control = pd.DataFrame([
        ("Filas", len(bruto), len(limpio_full), "No deben desaparecer registros"),
        ("Columnas principales", len(COLUMNAS), len(COLUMNAS), "10 sensores y contador smp"),
        ("Filas con datos en columnas adicionales", filas_extra_original, filas_extra_limpio, "El original presenta desalineaciones"),
        ("Celdas nulas en las 11 columnas", int(original.isna().sum().sum()), int(limpio.isna().sum().sum()), "Los nulos iniciales fueron tratados"),
    ], columns=["Indicador", "Antes", "Después", "Interpretación"])

    tabla_metodo = pd.DataFrame([
        ("Fuente antes", "Excel original, primeras 11 columnas principales sin corregir."),
        ("Fuente después", "Excel filtrado, hoja Datos; nulos y estructura corregidos."),
        ("Observaciones", "Se calculan los estadísticos con los datos numéricos disponibles por sensor."),
        ("DS", "Desviación estándar muestral (divisor N−1)."),
        ("IQR", "Rango intercuartílico = Q3 − Q1."),
        ("Límites IQR", "LI = Q1 − 1,5×IQR; LS = Q3 + 1,5×IQR."),
        ("Atípicos", "Valores por debajo de LI o por encima de LS. Se cuentan, no se eliminan."),
        ("Atípicos (%)", "100 × cantidad de atípicos / cantidad de valores no nulos del sensor."),
        ("Reducción DS (%)", "100 × (DS antes − DS después) / DS antes; puede ser negativa si la DS aumenta."),
        ("smp", "Contador de muestras: se incluye en nulos y control, no en atípicos de sensores."),
        ("Precaución", "Los atípicos del original pueden provenir de desalineación; no confirman fallas físicas."),
        ("Tratamiento pendiente", "La limpieza inicial no elimina atípicos. Aún debe justificarse su conservación o exclusión."),
    ], columns=["Concepto", "Criterio / explicación"])

    return {
        "Resumen": resumen,
        "Estadisticos": tabla_est,
        "Nulos": tabla_nulos,
        "Atipicos IQR": tabla_atipicos,
        "Control limpieza": tabla_control,
        "Metodologia": tabla_metodo,
    }


def escribir_excel(tablas, ruta_salida):
    """Crea hojas simples: sin filtros, sin paneles fijos y sin tablas dinámicas."""
    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as escritor:
        for nombre, tabla in tablas.items():
            tabla.to_excel(escritor, sheet_name=nombre, index=False)
            hoja = escritor.sheets[nombre]

            # TABLAS NORMALES: no se añaden filtros ni tablas de Excel.
            hoja.auto_filter.ref = None
            hoja.freeze_panes = None
            hoja.sheet_view.showGridLines = True
            hoja.row_dimensions[1].height = 28

            borde = Border(bottom=Side(style="thin", color="D9D9D9"))
            for celda in hoja[1]:
                celda.font = Font(name="Calibri", size=11, bold=True, color="262626")
                celda.fill = PatternFill(fill_type="solid", fgColor="E6E6E6")
                celda.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
                celda.border = borde

            for columna in hoja.columns:
                titulo = str(columna[0].value)
                if titulo == "Criterio / explicación":
                    ancho = 72
                elif titulo == "Interpretación":
                    ancho = 52
                elif titulo in {"Variable", "Indicador", "Concepto"}:
                    ancho = 40
                elif titulo in {"Sensor", "Etapa"}:
                    ancho = 20
                else:
                    ancho = 23

                letra = get_column_letter(columna[0].column)
                hoja.column_dimensions[letra].width = ancho
                for celda in columna[1:]:
                    celda.font = Font(name="Calibri", size=10, color="262626")
                    celda.alignment = Alignment(
                        vertical="center",
                        wrap_text=nombre in {"Metodologia", "Control limpieza"},
                    )
                    # Valores ya calculados por Python; se preservan los decimales.
                    if isinstance(celda.value, (int, float, np.number)):
                        if "(%)" in titulo or titulo == "Reducción DS (%)":
                            celda.number_format = '0.00"%"'
                        elif titulo in {"N", "Nulos antes", "Nulos después",
                                        "Atípicos antes", "Atípicos después",
                                        "Atípicos IQR", "Antes", "Después"}:
                            celda.number_format = "#,##0"
                        else:
                            celda.number_format = "0.000000"

            if nombre in {"Metodologia", "Control limpieza"}:
                for fila in range(2, hoja.max_row + 1):
                    hoja.row_dimensions[fila].height = 32

    print(f"Excel simple generado: {ruta_salida}")


def main():
    parser = argparse.ArgumentParser(description="Tablas comparativas original vs limpieza inicial")
    parser.add_argument("--original", help="Nombre/ruta del Excel original")
    parser.add_argument("--limpio", help="Nombre/ruta del Excel filtrado")
    parser.add_argument("--salida", help="Ruta del Excel de resultados (opcional)")
    args = parser.parse_args()
    carpeta = Path(__file__).resolve().parent

    original = localizar_archivo(carpeta, "CF-2023-I-204-1230_DAT*.xlsx", args.original, "original")
    limpio = localizar_archivo(carpeta, "Dataset_eolico_filtrado_2*.xlsx", args.limpio, "limpio")
    salida = (Path(args.salida).expanduser() if args.salida else carpeta / "Tablas_comparativas_Tarea2_SIMPLES.xlsx")
    if not salida.is_absolute():
        salida = carpeta / salida
    salida = salida.resolve()
    if salida in {original, limpio}:
        raise ValueError("La ruta de salida debe ser distinta de las entradas.")
    salida.parent.mkdir(parents=True, exist_ok=True)

    print(f"Original: {original.name}\nFiltrado: {limpio.name}", flush=True)
    bruto, datos_originales = leer_datos(original)
    limpio_full, datos_limpios = leer_datos(limpio, limpio=True)
    if len(datos_originales) != len(datos_limpios):
        print("AVISO: la cantidad de filas cambió entre original y filtrado.")
    tablas = preparar_tablas(bruto, datos_originales, limpio_full, datos_limpios)
    escribir_excel(tablas, salida)
    print(f"Hojas: {', '.join(tablas)}")
    print("Los Excel de entrada NO fueron modificados.")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError, PermissionError) as e:
        print(f"ERROR: {e}")
        raise SystemExit(1)
